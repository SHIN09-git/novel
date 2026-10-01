from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = "2E74B5"
ACCENT_DARK = "1F4D78"
TEXT = "263238"
MUTED = "66727A"
LIGHT_FILL = "E8EEF5"
CODE_FILL = "F4F6F8"
WARNING_FILL = "FFF4D6"
BODY_FONT = "Calibri"
EAST_ASIA_FONT = "Microsoft YaHei"
MONO_FONT = "Cascadia Mono"
TABLE_WIDTH_DXA = 9360


def set_run_font(run, *, size=None, bold=None, color=None, font=BODY_FONT):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT if font == BODY_FONT else font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def create_decimal_numbering(doc, start_value=1):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    nsid = OxmlElement("w:nsid")
    nsid.set(qn("w:val"), f"{abstract_id + 1:08X}")
    abstract.append(nsid)
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)
    template = OxmlElement("w:tmpl")
    template.set(qn("w:val"), f"{(abstract_id + 1) * 2654435761 & 0xFFFFFFFF:08X}")
    abstract.append(template)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), str(start_value))
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(start)
    level.append(num_fmt)
    level.append(level_text)
    level.append(level_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(tabs)
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start_value))
    level_override.append(start_override)
    num.append(level_override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = table._tbl.tblGrid
    table._tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    table._tbl.insert(1, grid)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def table_widths(rows):
    column_count = len(rows[0])
    if column_count == 1:
        return [TABLE_WIDTH_DXA]
    if column_count == 5:
        # Provider setup tables need enough room to keep service/provider labels
        # readable while leaving the URL and setup guidance columns dominant.
        return [1050, 1300, 2200, 2450, 2360]
    weights = []
    for column_index in range(column_count):
        longest = max(len(row[column_index]) if column_index < len(row) else 0 for row in rows)
        weights.append(max(8, min(longest, 42)))
    total = sum(weights)
    widths = [max(900, round(TABLE_WIDTH_DXA * weight / total)) for weight in weights]
    delta = TABLE_WIDTH_DXA - sum(widths)
    widths[-1] += delta
    if widths[-1] < 900:
        shortage = 900 - widths[-1]
        donor = max(range(len(widths) - 1), key=widths.__getitem__)
        widths[donor] -= shortage
        widths[-1] = 900
    return widths


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline_markdown(paragraph, text, *, size=11, color=TEXT):
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size - 0.5, color=ACCENT_DARK, font=MONO_FONT)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            set_run_font(run, size=size, color=ACCENT)
            run.underline = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def shade_paragraph(paragraph, fill, left_border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if left_border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "20")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_code_block(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.18)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.05
    shade_paragraph(paragraph, CODE_FILL, ACCENT)
    run = paragraph.add_run("\n".join(lines).rstrip())
    set_run_font(run, size=9.2, color=TEXT, font=MONO_FONT)


def add_markdown_table(doc, rows):
    column_count = max(len(row) for row in rows)
    normalized = [row + [""] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=column_count)
    table.style = "Table Grid"
    set_table_geometry(table, table_widths(normalized))
    for row_index, row in enumerate(normalized):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            add_inline_markdown(paragraph, value, size=9.2 if column_count >= 4 else 9.6)
            if row_index == 0:
                paragraph.paragraph_format.keep_with_next = True
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(ACCENT_DARK)
                set_cell_shading(cell, LIGHT_FILL)
    for row in table.rows:
        set_row_cant_split(row)
    set_repeat_table_header(table.rows[0])
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Novel Director 0.1.0  |  ")
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    number = OxmlElement("w:t")
    number.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, number, end])


def set_document_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, ACCENT, 18, 10),
        "Heading 2": (13, ACCENT, 14, 7),
        "Heading 3": (12, ACCENT_DARK, 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_sections(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("NOVEL DIRECTOR  /  新手操作手册")
    set_run_font(run, size=8.5, bold=True, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(44)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.LEFT
    kicker.paragraph_format.space_after = Pt(10)
    run = kicker.add_run("NOVEL DIRECTOR")
    set_run_font(run, size=10, bold=True, color=ACCENT)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("快速开始")
    set_run_font(run, size=30, bold=True, color=TEXT)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run("从 API 配置到第一章修订的完整新手指南")
    set_run_font(run, size=15, color=ACCENT_DARK)

    lead = doc.add_paragraph()
    lead.paragraph_format.left_indent = Inches(0.18)
    lead.paragraph_format.right_indent = Inches(0.18)
    lead.paragraph_format.space_before = Pt(6)
    lead.paragraph_format.space_after = Pt(18)
    lead.paragraph_format.line_spacing = 1.25
    shade_paragraph(lead, LIGHT_FILL, ACCENT)
    add_inline_markdown(
        lead,
        "这不是按钮清单。你将按真实工作流完成：安全配置 AI、准备最小设定、生成草稿、阅读诊断、选择修订方式、接受为正式章节并保留版本链。",
        size=11.5,
    )

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(4)
    run = meta.add_run("版本 0.1.0  |  中文优先  |  本地优先")
    set_run_font(run, size=10, bold=True, color=MUTED)

    privacy = doc.add_paragraph()
    privacy.paragraph_format.space_after = Pt(0)
    run = privacy.add_run("示例均为虚构数据。不要公开真实 API Key、私有稿件或项目 JSON。")
    set_run_font(run, size=9.5, color=MUTED)
    doc.add_page_break()


def parse_table(lines, start):
    if start + 1 >= len(lines) or not lines[start].lstrip().startswith("|"):
        return None
    separator = lines[start + 1].strip()
    if not re.match(r"^\|?\s*:?-{3,}", separator):
        return None
    rows = []
    index = start
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        raw = lines[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if index != start + 1:
            rows.append(cells)
        index += 1
    return rows, index


def add_body_from_markdown(doc, source):
    lines = source.splitlines()
    index = 0
    in_code = False
    code_lines = []
    first_h1_seen = False
    active_numbering_id = None

    while index < len(lines):
        line = lines[index].rstrip()

        if line.startswith("```"):
            active_numbering_id = None
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        table = parse_table(lines, index)
        if table:
            active_numbering_id = None
            rows, index = table
            add_markdown_table(doc, rows)
            continue

        if not line.strip():
            active_numbering_id = None
            index += 1
            continue

        if line.strip() == "---":
            active_numbering_id = None
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(2)
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            active_numbering_id = None
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level == 1 and not first_h1_seen:
                first_h1_seen = True
                index += 1
                continue
            if title == "English":
                doc.add_page_break()
            style = "Heading 1" if level <= 2 else "Heading 2" if level == 3 else "Heading 3"
            paragraph = doc.add_paragraph(style=style)
            add_inline_markdown(paragraph, title, size={"Heading 1": 16, "Heading 2": 13, "Heading 3": 12}[style], color=ACCENT if style != "Heading 3" else ACCENT_DARK)
            index += 1
            continue

        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        numbered = re.match(r"^\s*(\d+)\.\s+(.+)$", line)
        if bullet or numbered:
            if bullet:
                active_numbering_id = None
                paragraph = doc.add_paragraph(style="List Bullet")
                item_text = bullet.group(1)
            else:
                if active_numbering_id is None:
                    active_numbering_id = create_decimal_numbering(doc, int(numbered.group(1)))
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.25
                apply_numbering(paragraph, active_numbering_id)
                item_text = numbered.group(2)
            add_inline_markdown(paragraph, item_text)
            index += 1
            continue

        if line.startswith(">"):
            active_numbering_id = None
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            paragraph.paragraph_format.space_after = Pt(8)
            shade_paragraph(paragraph, WARNING_FILL, "D59B25")
            add_inline_markdown(paragraph, line.lstrip("> "))
            index += 1
            continue

        active_numbering_id = None
        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines):
            candidate = lines[index].rstrip()
            if not candidate.strip() or candidate.startswith(("#", "```", "|", ">")):
                break
            if re.match(r"^\s*[-*]\s+", candidate) or re.match(r"^\s*\d+\.\s+", candidate):
                break
            paragraph_lines.append(candidate.strip())
            index += 1
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.widow_control = True
        add_inline_markdown(paragraph, " ".join(paragraph_lines))


def build(source_path: Path, output_path: Path):
    source = source_path.read_text(encoding="utf-8")
    doc = Document()
    set_document_styles(doc)
    configure_sections(doc)
    add_cover(doc)
    add_body_from_markdown(doc, source)

    props = doc.core_properties
    props.title = "Novel Director 快速开始"
    props.subject = "从 AI API 配置到章节修订的完整新手指南"
    props.author = "Novel Director contributors"
    props.keywords = "Novel Director, Quickstart, AI writing, revision, API"
    props.comments = "Generated from QUICKSTART.md by scripts/build-quickstart-docx.py"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main():
    root = Path(__file__).resolve().parent.parent
    source = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else root / "QUICKSTART.md"
    output = Path(sys.argv[2]).resolve() if len(sys.argv) > 2 else root / "QUICKSTART.docx"
    build(source, output)
    print(f"Built {output}")


if __name__ == "__main__":
    main()
